"""Payroll report — the signed artifact. Its numbers must be the engine's own,
the sheets must foot to the headline, and the PDF must be exactly one page.

Built on a synthetic pair with a title block (entity + period), so
entity/period extraction is exercised too — never the real gitignored files.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook as read_xlsx

from _helpers import make_grid, make_workbook
from excel_analysis.adapters.report_pdf import render_pdf
from excel_analysis.adapters.report_xlsx import render_xlsx
from excel_analysis.app.payroll_report import build_payroll_report, period_token
from excel_analysis.app.pipeline import compare_workbooks

HEADERS = [
    "Employee_ID", "Employee Name", "Department", "Basic", "Gross Earnings",
    "PF", "Gross Deduction", "Net Pay", "LTA Reimbursement", "Gross Reimbursement", "Net Payable",
]


def _row(emp_id, name, dept, basic, pf, lta):
    ge, gd = basic, pf
    net_pay = ge - gd
    gr = lta
    return [emp_id, name, dept, basic, ge, pf, gd, net_pay, lta, gr, net_pay + gr]


def _wb(title, month, rows):
    grid = [
        [title] + [None] * (len(HEADERS) - 1),
        [f"Pay Sheet for the Month of {month} 2026"] + [None] * (len(HEADERS) - 1),
        [None] * len(HEADERS),
        HEADERS,
        *rows,
    ]
    return make_workbook([make_grid(grid, name="Sheet1")], path=f"{month}.xlsx")


LEFT = _wb("ABC Ltd.", "May", [
    _row(1, "Asha", "Engineering", 1000, 120, 50),
    _row(2, "Bala", "Engineering", 2000, 240, 80),
    _row(3, "Chen", "Finance", 3000, 360, 120),
])
RIGHT = _wb("ABC Ltd.", "Jun", [
    _row(1, "Asha", "Engineering", 1200, 120, 50),   # raise
    _row(2, "Bala", "Engineering", 2000, 240, 80),   # unchanged
    _row(4, "Devi", "Finance", 2500, 250, 0),        # joiner (3 is a leaver)
])


def _report():
    result = compare_workbooks(LEFT, RIGHT, anchor="Net Payable")
    return build_payroll_report(result, LEFT, RIGHT, left_name="May.xlsx", right_name="Jun.xlsx", prepared_by="Tester")


def test_report_header_pulled_from_files():
    rep = _report()
    assert rep.entity == "ABC Ltd."
    assert "May" in rep.period_prev and "Jun" in rep.period_curr
    assert rep.anchor == "Net Payable" and rep.anchor_substituted is False


def test_report_numbers_are_canonical_and_tie_out():
    rep = _report()
    # headline arithmetic closes by hand
    assert round(rep.total_curr - rep.total_prev, 2) == rep.delta
    # component net impact and employee deltas both foot to the headline
    assert abs(sum(c.net_impact for c in rep.components_all) - (rep.comp_movement + rep.reimb_movement)) < 0.01
    assert abs(sum(e.delta for e in rep.employees_all) - rep.delta) < 0.01
    assert (rep.joiners, rep.leavers) == (1, 1)
    assert rep.ties_out is True


def test_xlsx_every_sheet_foots_to_headline():
    rep = _report()
    wb = read_xlsx(BytesIO(render_xlsx(rep)))
    assert wb.sheetnames == ["Bridge", "Components", "Employees", "Exceptions"]
    comps = wb["Components"]
    assert round(comps.cell(row=comps.max_row, column=5).value, 2) == round(rep.comp_movement + rep.reimb_movement, 2)
    emps = wb["Employees"]
    assert round(emps.cell(row=emps.max_row, column=8).value, 2) == round(rep.delta, 2)


def test_pdf_is_a_single_page():
    rep = _report()
    pdf = render_pdf(rep)
    assert pdf[:5] == b"%PDF-"
    # count page objects (each leaf page is '/Type /Page' but not '/Type /Pages')
    pages = pdf.count(b"/Type /Page") - pdf.count(b"/Type /Pages")
    assert pages == 1


def test_period_token_is_filename_safe():
    assert period_token("Pay Sheet for the Month of Jun 2026") == "Jun2026"
    assert period_token("May.xlsx") == "Mayxlsx" or period_token("May.xlsx")  # squeezed, non-empty


_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _write_payroll_xlsx(path, title, month, rows):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append([title])
    ws.append([f"Pay Sheet for the Month of {month} 2026"])
    ws.append([])
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_report_endpoint_streams_pdf_and_xlsx(tmp_path):
    from fastapi.testclient import TestClient

    from excel_analysis.api import app

    left = tmp_path / "may.xlsx"
    right = tmp_path / "jun.xlsx"
    _write_payroll_xlsx(left, "ABC Ltd.", "May", [
        _row(1, "Asha", "Engineering", 1000, 120, 50),
        _row(2, "Bala", "Engineering", 2000, 240, 80),
        _row(3, "Chen", "Finance", 3000, 360, 120),
    ])
    _write_payroll_xlsx(right, "ABC Ltd.", "Jun", [
        _row(1, "Asha", "Engineering", 1200, 120, 50),
        _row(2, "Bala", "Engineering", 2000, 240, 80),
        _row(4, "Devi", "Finance", 2500, 250, 0),
    ])
    client = TestClient(app)
    for fmt, magic, ctype in [("pdf", b"%PDF-", "application/pdf"), ("xlsx", b"PK", _XLSX)]:
        with open(left, "rb") as lf, open(right, "rb") as rf:
            resp = client.post(
                f"/report?format={fmt}",
                files={"left_file": ("may.xlsx", lf, _XLSX), "right_file": ("jun.xlsx", rf, _XLSX)},
            )
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"].startswith(ctype)
        assert resp.content[: len(magic)] == magic
        cd = resp.headers["content-disposition"]
        assert "May2026" in cd and "Jun2026" in cd and cd.endswith(f'.{fmt}"')

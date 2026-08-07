"""Deterministic builders for the messy-workbook fixture corpus (M1).

Each builder constructs one workbook that exercises a specific edge case. Inputs
are built in code (not committed binaries) so they are stable by construction and
reviewable in diffs; the committed golden snapshots under tests/golden/ are the
oracle. The real Payroll Analytics.xlsx is a review artifact, never a test
oracle (it holds real personal data).
"""

from __future__ import annotations

import os
from datetime import datetime

import openpyxl


def _simple(wb):
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Age", "City"])
    ws.append(["Alice", 30, "London"])
    ws.append(["Bob", 25, "Paris"])
    ws.append(["Carol", 41, "Berlin"])
    ws.append(["Bob", 25, "Paris"])  # duplicate row -> unique_count < count


def _title_and_blank(wb):
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "Quarterly Report"  # merged title above the table
    ws.merge_cells("A1:C1")
    # row 2 intentionally blank
    ws["A3"] = "Region"
    ws["B3"] = "Units"
    ws["C3"] = "Revenue"
    for i, (reg, u, rev) in enumerate(
        [("North", 120, 1500.5), ("South", 95, 1120.0), ("East", 60, 700.0)], start=4
    ):
        ws.cell(i, 1, reg)
        ws.cell(i, 2, u)
        ws.cell(i, 3, rev)


def _stacked_tables(wb):
    ws = wb.active
    ws.title = "Two"
    ws["A1"] = "Product"
    ws["B1"] = "Price"
    ws["A2"] = "Widget"
    ws["B2"] = 9.99
    ws["A3"] = "Gadget"
    ws["B3"] = 19.99
    # row 4 blank -> second stacked table
    ws["A5"] = "Region"
    ws["B5"] = "Total"
    ws["A6"] = "North"
    ws["B6"] = 100
    ws["A7"] = "South"
    ws["B7"] = 200


def _merged_header(wb):
    ws = wb.active
    ws.title = "Merged"
    ws["A1"] = "Location"  # merged across two columns
    ws.merge_cells("A1:B1")
    ws["C1"] = "Amount"
    ws["A2"] = "London"
    ws["B2"] = "UK"
    ws["C2"] = 10
    ws["A3"] = "Paris"
    ws["B3"] = "FR"
    ws["C3"] = 20


def _mixed_types(wb):
    ws = wb.active
    ws.title = "Mixed"
    ws["A1"] = "Item"
    ws["B1"] = "Qty"
    ws["A2"] = "Apple"
    ws["B2"] = 10
    ws["A3"] = "Pear"
    ws["B3"] = "N/A"  # text in an otherwise-numeric column
    ws["A4"] = "Plum"
    ws["B4"] = 5


def _hidden_sheet(wb):
    ws = wb.active
    ws.title = "Main"
    ws.append(["Code", "Value"])
    ws.append(["A", 1])
    ws.append(["B", 2])
    notes = wb.create_sheet("Notes")
    notes["A1"] = "internal"
    notes["A2"] = "draft"
    notes.sheet_state = "hidden"


def _duplicate_and_blank_headers(wb):
    ws = wb.active
    ws.title = "Dups"
    ws["A1"] = "ID"
    ws["B1"] = "Val"
    ws["C1"] = "Val"  # duplicate header
    # D1 blank header, but the column has data below
    ws["A2"] = 1
    ws["B2"] = 10
    ws["C2"] = 100
    ws["D2"] = "x"
    ws["A3"] = 2
    ws["B3"] = 20
    ws["C3"] = 200
    ws["D3"] = "y"


def _empty_sheet(wb):
    ws = wb.active
    ws.title = "Empty"
    # no cells written -> detector should report no tables, not crash


def _single_column(wb):
    ws = wb.active
    ws.title = "One"
    ws["A1"] = "Fruit"
    ws["A2"] = "Apple"
    ws["A3"] = "Banana"
    ws["A4"] = "Cherry"


def _dates_and_bools(wb):
    ws = wb.active
    ws.title = "Types"
    ws["A1"] = "When"
    ws["B1"] = "Active"
    ws["C1"] = "Count"
    for i, (d, a, c) in enumerate(
        [
            (datetime(2026, 1, 1), True, 5),
            (datetime(2026, 2, 1), False, 10),
            (datetime(2026, 3, 1), True, 15),
        ],
        start=2,
    ):
        ws.cell(i, 1, d)
        ws.cell(i, 2, a)
        ws.cell(i, 3, c)


def _no_header_numeric(wb):
    ws = wb.active
    ws.title = "NoHeader"
    # All-numeric with repeated values -> low header confidence (< 0.5).
    for i, row in enumerate([[1, 1, 1], [2, 2, 2], [3, 3, 3], [4, 4, 4]], start=1):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v)


def _unstructured_notes(wb):
    ws = wb.active
    ws.title = "Notes"
    # Single column of multi-word prose -> reported as unstructured, not profiled.
    ws["A1"] = "Internal draft - do not distribute"
    ws["A2"] = "Reviewed by finance on Tuesday"
    ws["A3"] = "Pending approval from the director"


FIXTURES = {
    "simple": _simple,
    "title_and_blank": _title_and_blank,
    "stacked_tables": _stacked_tables,
    "merged_header": _merged_header,
    "mixed_types": _mixed_types,
    "hidden_sheet": _hidden_sheet,
    "duplicate_and_blank_headers": _duplicate_and_blank_headers,
    "empty_sheet": _empty_sheet,
    "single_column": _single_column,
    "dates_and_bools": _dates_and_bools,
    "no_header_numeric": _no_header_numeric,
    "unstructured_notes": _unstructured_notes,
}


def build_all(dest_dir: str) -> dict[str, str]:
    """Build every fixture into dest_dir; return {name: path}."""
    os.makedirs(dest_dir, exist_ok=True)
    paths: dict[str, str] = {}
    for name, builder in FIXTURES.items():
        wb = openpyxl.Workbook()
        builder(wb)
        path = os.path.join(dest_dir, f"{name}.xlsx")
        wb.save(path)
        paths[name] = path
    return paths


# --- Synthetic compare pair (M6): a keyed diff exercising every field --------
# Mirrors the real file-1/file-2 shape (matched/added/removed/retyped columns,
# numeric + categorical changes, a quiet date<->text retype, added/removed rows)
# so a committed golden needs no gitignored real data.


def _compare_left(wb):
    ws = wb.active
    ws.title = "Pay"
    ws.append(["Emp_ID", "Name", "Salary", "City", "DoB", "ExitReason"])
    ws.append([1, "Alice", 100, "London", datetime(1990, 1, 1), ""])
    ws.append([2, "Bob", 200, "Paris", datetime(1991, 2, 2), ""])
    ws.append([3, "Carol", 300, "Rome", datetime(1992, 3, 3), ""])
    ws.append([4, "Dave", 400, "Berlin", datetime(1993, 4, 4), "resigned"])


def _compare_right(wb):
    ws = wb.active
    ws.title = "Pay"
    ws.append(["Emp_ID", "Name", "Salary", "City", "DoB", "Bonus"])
    ws.append([2, "Bob", 220, "Paris", "1991-02-02", 10])    # Salary +10%; DoB quiet
    ws.append([3, "Carol", 330, "Turin", "1991-02-02", 20])  # Salary +10%; City change; DoB shifted (defect)
    ws.append([4, "Dave", 440, "Berlin", "1993-04-04", 30])  # Salary +10%; DoB quiet
    ws.append([5, "Eve", 550, "Madrid", "1994-05-05", 40])   # added row


def build_compare_pair(dest_dir: str) -> dict[str, str]:
    os.makedirs(dest_dir, exist_ok=True)
    paths: dict[str, str] = {}
    for name, builder in [("compare_left", _compare_left), ("compare_right", _compare_right)]:
        wb = openpyxl.Workbook()
        builder(wb)
        path = os.path.join(dest_dir, f"{name}.xlsx")
        wb.save(path)
        paths[name] = path
    return paths

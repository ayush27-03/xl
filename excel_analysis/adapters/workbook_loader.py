"""WorkbookLoader adapter — the ONLY module that imports openpyxl (CLAUDE.md
rule 2).

It translates an .xlsx file into the raw-grid domain models and does nothing
more: no header detection, no cleaning, no interpretation. Cell types are
inferred from the *value* Python type (with data_only=True openpyxl returns
computed values), which is robust and keeps openpyxl's own type codes out of the
domain.
"""

from __future__ import annotations

import os
import zipfile
from datetime import date, datetime

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from ..domain.errors import WorkbookLoadError
from ..domain.models import (
    CellType,
    MergedRange,
    RawCell,
    RawSheetGrid,
    RawWorkbook,
)

_EMPTY_CELL = RawCell(value=None, type=CellType.EMPTY)

# Literal Excel error strings that may surface as cached values.
_EXCEL_ERRORS = frozenset(
    {
        "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!",
        "#SPILL!", "#CALC!", "#GETTING_DATA",
    }
)


def load_workbook(path: str) -> RawWorkbook:
    """Load an .xlsx file into a RawWorkbook. Fatal on missing/invalid files."""
    if not os.path.isfile(path):
        raise WorkbookLoadError(f"File not found: {path}")
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except (InvalidFileException, zipfile.BadZipFile, OSError, KeyError, ValueError) as exc:
        raise WorkbookLoadError(
            f"Could not open '{path}' as an .xlsx workbook: {exc}"
        ) from exc

    try:
        sheets = tuple(_read_sheet(ws) for ws in wb.worksheets)
    finally:
        wb.close()
    return RawWorkbook(path=os.path.abspath(path), sheets=sheets)


def _read_sheet(ws) -> RawSheetGrid:
    hidden = ws.sheet_state != "visible"  # 'hidden' or 'veryHidden'
    n_rows = ws.max_row or 0
    n_cols = ws.max_column or 0

    if n_rows == 0 or n_cols == 0:
        return RawSheetGrid(
            name=ws.title, hidden=hidden, n_rows=0, n_cols=0,
            cells=(), merged_ranges=(),
        )

    rows: list[tuple[RawCell, ...]] = []
    for row in ws.iter_rows(min_row=1, max_row=n_rows, min_col=1, max_col=n_cols):
        cells = [_to_raw_cell(c.value) for c in row]
        if len(cells) < n_cols:  # defensive: pad short rows to a rectangle
            cells.extend([_EMPTY_CELL] * (n_cols - len(cells)))
        rows.append(tuple(cells))

    merged = tuple(
        MergedRange(
            first_row=r.min_row - 1, last_row=r.max_row - 1,
            first_col=r.min_col - 1, last_col=r.max_col - 1,
        )
        for r in ws.merged_cells.ranges
    )
    return RawSheetGrid(
        name=ws.title, hidden=hidden, n_rows=n_rows, n_cols=n_cols,
        cells=tuple(rows), merged_ranges=merged,
    )


def _to_raw_cell(value) -> RawCell:
    if value is None:
        return _EMPTY_CELL
    # bool must precede int/float: bool is a subclass of int.
    if isinstance(value, bool):
        return RawCell(value=value, type=CellType.BOOL)
    if isinstance(value, (datetime, date)):
        return RawCell(value=value, type=CellType.DATE)
    if isinstance(value, (int, float)):
        return RawCell(value=value, type=CellType.NUMBER)
    if isinstance(value, str):
        stripped = value.strip()
        if stripped == "":
            return _EMPTY_CELL
        if stripped in _EXCEL_ERRORS:
            return RawCell(value=value, type=CellType.ERROR)
        return RawCell(value=value, type=CellType.TEXT)
    return RawCell(value=value, type=CellType.TEXT)
